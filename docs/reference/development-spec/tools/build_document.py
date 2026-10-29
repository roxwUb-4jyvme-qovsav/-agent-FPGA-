"""Render the specification source to an editable Word document."""
from pathlib import Path
import datetime, hashlib, json, re, shutil
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'FPGA多Agent协同开发平台开发规格书_v1.0.md'
OUT=SOURCE.with_suffix('.docx')
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','Caption']:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    st.paragraph_format.line_spacing=1.15;st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.widow_control=True
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(18)
doc.styles['Subtitle'].font.size=Pt(12);doc.styles['Subtitle'].font.italic=False
for name,size in [('Heading 1',15),('Heading 2',11.5),('Heading 3',11)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
def inline(p,text):
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text):
        if part.startswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        elif part.startswith('**'):p.add_run(part[2:-2]).bold=True
        else:p.add_run(part)
    return p
def para(text,style=None):return inline(doc.add_paragraph(style=style),text)
def table(rows,widths=None):
    count=len(rows[0]);total=17.2
    if widths is None:
        widths={2:[4.3,12.9],3:[3.7,6,7.5],4:[3.1,4.4,4.8,4.9],5:[3.7,2.7,3.2,3.8,3.8]}[count]
        if rows[0][0]=='偏移':widths=[2,2,4,9.2]
        if rows[0][0]=='动作':widths=[6.1,2.3,2.8,3,3]
        if rows[0][0]=='类型':widths=[2.5,6.6,8.1]
        if rows[0][0]=='错误码':widths=[3.4,6.3,7.5]
        if rows[0][0]=='需求组':widths=[3.2,10.3,3.7]
        if rows[0][0]=='目录':widths=[4.6,6.3,6.3]
        if rows[0][0]=='数据表组':widths=[5.1,5.8,6.3]
        if rows[0][0]=='HTTP 状态':widths=[2.5,7.5,7.2]
        if rows[0][0]=='模块':widths=[3.6,7,6.6]
        if rows[0][0]=='参数':widths=[3.2,5.2,8.8]
    t=doc.add_table(rows=0,cols=count);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for n,values in enumerate(rows):
        row=t.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if n==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=inline(cell.paragraphs[0],value);p.paragraph_format.line_spacing=1.08
            p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            if n==0:p.paragraph_format.keep_with_next=True
            for r in p.runs:r.font.size=Pt(9.3);r.bold=n==0
            mar=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            cell._tc.get_or_add_tcPr().append(mar)
            if n==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE7EF');cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=Pt(2)
    return t
text=SOURCE.read_text(encoding='utf-8');lines=text.splitlines();i=0;toc=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('## 1 ') and not toc:
        doc.add_page_break();para('目录','Heading 1')
        # Only numbered chapters and annexes enter the TOC; the TOC title does not.
        doc.paragraphs[-1].style=doc.styles['Normal'];doc.paragraphs[-1].runs[0].font.size=Pt(16);doc.paragraphs[-1].runs[0].bold=True
        p=doc.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');p._p.append(f)
        doc.add_page_break();toc=True
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',x) for x in cells):rows.append(cells)
            i+=1
        table(rows);continue
    if line.startswith('# '):para(line[2:],'Title')
    elif line.startswith('## '):para(line[3:],'Heading 1')
    elif line.startswith('### '):para(line[4:],'Heading 2')
    elif line.startswith('- '):para(line[2:],'List Bullet')
    else:para(line)
    i+=1
doc.add_page_break();para('附录 A 验收分组与检查项','Heading 1')
para('以下为验收计划索引，全部条目均为 P0。实际操作、预期与证据要求见随附 acceptance/catalog.json；NOT_RUN 表示本规格未执行产品验收。')
cases=json.loads((ROOT/'acceptance/catalog.json').read_text(encoding='utf-8'))['tests']
rows=[['验收标识','检查项','预期结果']]+[[x['id'],x['title'],x['expected']] for x in cases]
table(rows,[2.8,4.0,10.4])
para('附录 B 工作包与交接清单','Heading 1')
plan=json.loads((ROOT/'planning/backlog.json').read_text(encoding='utf-8'))
para('总估算为 29 至 49 个编码人日。以下任务按依赖领取，完整上下游与状态记录见 planning/backlog.json。')
rows=[['任务与责任','交付内容','前置任务']]
for wp in plan['work_packages']:
    for x in wp['tasks']:rows.append([x['id']+'  '+x['owner'],x['deliverable'],'、'.join(x['depends_on']) or '无'])
table(rows,[3,10.8,3.4])
hp=sec.header.paragraphs[0];hp.text='FPGA 多 Agent 协同开发平台　开发规格书 1.0'
for r in hp.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('2026 10 29　　第 ')
for name,after in [('PAGE',' 页　共 '),('NUMPAGES',' 页')]:
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),name);fp._p.append(f);fp.add_run(after)
for r in fp.runs:r.font.size=Pt(8)
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
doc.core_properties.title='FPGA 多 Agent 协同开发平台开发规格书';doc.core_properties.author=''
doc.core_properties.subject='本机工作台 统一契约 采集协议 独立验证 完整管理生态'
if OUT.exists():
    back=ROOT/'内部校验'/'备份'/datetime.datetime.now().strftime('%Y%m%d%H%M%S%f')/OUT.name
    back.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(OUT,back)
tmp=OUT.with_suffix('.tmp.docx');doc.save(tmp);tmp.replace(OUT)
print(json.dumps({'output':str(OUT),'characters':len(text),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
