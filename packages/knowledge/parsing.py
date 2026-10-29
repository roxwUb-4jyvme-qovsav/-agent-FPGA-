"""本地文本提取。图片、扫描页和不支持的媒体返回明确阻断。"""
import csv
import hashlib
import zipfile
from io import BytesIO, StringIO
from typing import Any


def _chunks(text: str, prefix: str = "", max_chars: int = 2400) -> list[dict[str, Any]]:
    lines = text.splitlines()
    result: list[dict[str, Any]] = []
    start, part, size = 1, [], 0
    for number, line in enumerate(lines, 1):
        if part and size + len(line) > max_chars:
            value = "\n".join(part)
            result.append({"text": value, "locator": f"{prefix}lines:{start}-{number - 1}", "content_sha256": hashlib.sha256(value.encode()).hexdigest()})
            start, part, size = number, [], 0
        part.append(line)
        size += len(line) + 1
    if part and any(line.strip() for line in part):
        value = "\n".join(part)
        result.append({"text": value, "locator": f"{prefix}lines:{start}-{len(lines)}", "content_sha256": hashlib.sha256(value.encode()).hexdigest()})
    return result


def _zip_size_guard(content: bytes) -> None:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("DOCUMENT_EXPANSION_LIMIT")


def parse_document(content: bytes, media_type: str, *, source_uri: str) -> dict[str, Any]:
    result: dict[str, Any] = {"status": "BLOCKED", "source_uri": source_uri, "content_sha256": hashlib.sha256(content).hexdigest(), "chunks": [], "diagnostics": []}
    if len(content) > 50 * 1024 * 1024:
        result["diagnostics"] = ["DOCUMENT_SIZE_LIMIT"]
        return result
    kind = media_type.split(";", 1)[0].strip().lower()
    try:
        if kind.startswith("text/") or kind in {"application/json", "application/xml"}:
            result["chunks"] = _chunks(content.decode("utf-8-sig"))
        elif kind == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                result["diagnostics"] = ["PDF_ENCRYPTED"]
                return result
            empty_pages: list[int] = []
            for page_number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    empty_pages.append(page_number)
                result["chunks"].extend(_chunks(text, f"page:{page_number};"))
            if empty_pages:
                result["diagnostics"].append("OCR_REQUIRED_PAGES:" + ",".join(map(str, empty_pages)))
                # 部分页未解析时不声称完整解析成功，已提取片段仍留给审阅。
                return result
        elif kind == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            from docx import Document
            _zip_size_guard(content)
            document = Document(BytesIO(content))
            for number, paragraph in enumerate(document.paragraphs, 1):
                result["chunks"].extend(_chunks(paragraph.text, f"paragraph:{number};"))
            for table_number, table in enumerate(document.tables, 1):
                header = " | ".join(cell.text for cell in table.rows[0].cells) if table.rows else ""
                for row_number, row in enumerate(table.rows[1:], 2):
                    result["chunks"].extend(_chunks(header + "\n" + " | ".join(cell.text for cell in row.cells), f"table:{table_number};row:{row_number};"))
        elif kind == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            from openpyxl import load_workbook
            _zip_size_guard(content)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, ())
                for row_number, row in enumerate(rows, 2):
                    stream = StringIO()
                    writer = csv.writer(stream)
                    writer.writerow(header)
                    writer.writerow(row)
                    result["chunks"].extend(_chunks(stream.getvalue(), f"sheet:{sheet.title};row:{row_number};"))
            workbook.close()
        else:
            result["diagnostics"] = ["OCR_NOT_CONFIGURED" if kind.startswith("image/") else "MEDIA_TYPE_UNSUPPORTED"]
            return result
    except ImportError:
        result["diagnostics"] = ["PARSER_DEPENDENCY_UNAVAILABLE"]
        return result
    except Exception as exc:
        # 解析错误只保存错误种类，不能把附件内容或主机路径带入诊断。
        result.update(status="FAILED", diagnostics=[f"PARSE_FAILED:{type(exc).__name__}"])
        return result
    result["status"] = "PARSED" if result["chunks"] else "BLOCKED"
    if not result["chunks"]:
        result["diagnostics"].append("NO_EXTRACTABLE_TEXT")
    return result
