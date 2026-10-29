from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import uuid4

import httpx
import pytest

from packages.knowledge.index import KnowledgeError, PgvectorSearch
from packages.knowledge.parsing import parse_document
from packages.model_gateway.budget import BudgetError, BudgetLedger
from packages.orchestration.dag import can_reuse
from packages.quotes.pricing import QuoteError, selling_price
from packages.tool_gateway.extensions import ExtensionError, transition_extension
from packages.tool_gateway.mcp import ControlledMCPClient, MCPError, tools_schema_hash
from packages.tool_gateway.sdk import SDKError, validate_adapter_response
from workers.health import detect_tools


def test_budget_requires_all_four_scope_levels_and_json_roundtrip():
    ledger = BudgetLedger({"global": "1"})
    with pytest.raises(BudgetError, match="SCOPES"):
        ledger.reserve("inv", ["global"], "0.1", "v1")
    copy = BudgetLedger(**ledger.to_dict())
    assert copy.to_dict() == ledger.to_dict()


def test_reuse_rejects_missing_hashes_and_revoked_authority():
    fingerprints = {key: "a" for key in ("artifact_sha256", "test_sha256", "toolchain_sha256", "generator_sha256", "config_sha256", "input_sha256")}
    fingerprints["authorized"] = True
    assert can_reuse(fingerprints, deepcopy(fingerprints))
    assert not can_reuse(fingerprints, {**fingerprints, "authorized": False})
    assert not can_reuse({"authorized": True}, {"authorized": True})


def test_docx_and_xlsx_tables_keep_headers_units_and_locations():
    from docx import Document
    from openpyxl import Workbook
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "型号", "电压/V"
    table.cell(1, 0).text, table.cell(1, 1).text = "XC7A35T", "3.3"
    stream = BytesIO()
    document.save(stream)
    parsed = parse_document(stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", source_uri="artifact:docx")
    assert parsed["status"] == "PARSED"
    assert "table:1;row:2;" in parsed["chunks"][0]["locator"]
    assert "电压/V" in parsed["chunks"][0]["text"]
    book = Workbook()
    book.active.append(["Part", "V"])
    book.active.append(["XC7A35T", 3.3])
    stream = BytesIO()
    book.save(stream)
    parsed = parse_document(stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", source_uri="artifact:xlsx")
    assert parsed["status"] == "PARSED"
    assert "sheet:Sheet;row:2;" in parsed["chunks"][0]["locator"]
    assert "Part,V" in parsed["chunks"][0]["text"]


def test_pgvector_pre_filters_authorized_documents_and_generation_before_ranking():
    observed = []
    class Connection:
        def execute(self, sql, parameters):
            observed.append((sql, parameters))
            return self
        def fetchall(self):
            return []
    adapter = PgvectorSearch(Connection(), generation="g2", model_version="embed-v2", dimensions=2)
    assert adapter.search([0.1, 0.2], project_id="p", authorized_document_ids=["d"]) == []
    assert "document_id = ANY" in observed[0][0]
    assert observed[0][1][2:6] == (["d"], "g2", "embed-v2", 2)
    with pytest.raises(KnowledgeError):
        adapter.search([0.1], project_id="p", authorized_document_ids=["d"])
    assert adapter.search([0.1, 0.2], project_id="p", authorized_document_ids=[]) == []
    assert len(observed) == 1


def tools():
    return [{"name": "lint", "inputSchema": {"type": "object", "properties": {"artifact_id": {"type": "string"}}, "required": ["artifact_id"], "additionalProperties": False}}]


def config():
    return {"transport": "streamable_http", "endpoint": "https://mcp.test/mcp", "timeout_seconds": 1, "schema_sha256": tools_schema_hash(tools())}


def context():
    return {"manage_connections": True, "member": True, "project_id": "p", "revision_id": "r", "enabled": True, "trusted": True, "healthy": True, "capabilities": ["rtl.lint"]}


def invocation():
    return {"project_id": "p", "revision_id": "r", "capability": "rtl.lint", "input": {"artifact_id": "a"}, "deadline": (datetime.now(timezone.utc) + timedelta(seconds=10)).isoformat(), "approval_ids": []}


def test_mcp_http_initialization_discovery_schema_lock_and_controlled_call():
    observed = []
    def handler(req):
        import json
        data = json.loads(req.content)
        observed.append(data["method"])
        results = {"initialize": {"protocolVersion": "2025-03-26", "capabilities": {"tools": {}}, "serverInfo": {"name": "test", "version": "1"}},
                   "tools/list": {"tools": tools()}, "tools/call": {"content": [{"type": "text", "text": "done"}], "isError": False}}
        if data["method"].startswith("notifications/"):
            return httpx.Response(202)
        return httpx.Response(200, json={"jsonrpc": "2.0", "id": data["id"], "result": results[data["method"]]}, headers={"Mcp-Session-Id": "session"})
    client = ControlledMCPClient(config(), capability_tools={"rtl.lint": "lint"}, allowed_hosts=["mcp.test"], resolver=lambda host: ["8.8.8.8"], transport=httpx.MockTransport(handler))
    assert client.connect(context())["tools"] == tools()
    assert client.call(invocation(), context())["content"][0]["text"] == "done"
    assert observed == ["initialize", "notifications/initialized", "tools/list", "tools/call"]
    with pytest.raises(MCPError, match="INPUT"):
        client.call({**invocation(), "input": {}}, context())
    with pytest.raises(ValueError):
        client.call(invocation(), {**context(), "trusted": False})


def test_mcp_schema_changes_and_unmanaged_stdio_are_blocked():
    client = ControlledMCPClient({**config(), "transport": "stdio"}, capability_tools={})
    with pytest.raises(MCPError, match="STDIO"):
        client.connect(context())
    def handler(req):
        import json
        data = json.loads(req.content)
        if data["method"].startswith("notifications/"):
            return httpx.Response(202)
        result = {"protocolVersion": "2025-03-26"} if data["method"] == "initialize" else {"tools": []}
        return httpx.Response(200, json={"jsonrpc": "2.0", "id": data["id"], "result": result})
    client = ControlledMCPClient(config(), capability_tools={}, allowed_hosts=["mcp.test"], resolver=lambda h: ["8.8.8.8"], transport=httpx.MockTransport(handler))
    with pytest.raises(MCPError, match="SCHEMA_CHANGED"):
        client.connect(context())


def test_sdk_rejects_wrong_identity_schema_and_success_without_payload():
    identifier = str(uuid4())
    result = {"schema_version": "1.0", "invocation_id": identifier, "status": "completed", "output_schema": "test", "output": {"ok": True}, "diagnostics": [], "request_id": None}
    validate_adapter_response(result, expected_invocation_id=identifier, expected_output_schema="test", output_validator=lambda name, value: None)
    with pytest.raises(SDKError):
        validate_adapter_response(result, expected_invocation_id=str(uuid4()), expected_output_schema="test", output_validator=lambda name, value: None)
    with pytest.raises(SDKError):
        validate_adapter_response({**result, "output": None}, expected_invocation_id=identifier, expected_output_schema="test", output_validator=lambda name, value: None)


def test_extension_enable_needs_current_health_and_approval():
    with pytest.raises(ExtensionError):
        transition_extension("DISABLED", "ENABLED", dependencies_valid=True, compatible=True, trusted=True)
    assert transition_extension("DISABLED", "ENABLED", dependencies_valid=True, compatible=True, trusted=True, healthy=True, approved=True) == "ENABLED"


def test_margin_cannot_equal_one_and_health_has_no_fake_hardware_result():
    with pytest.raises(QuoteError):
        selling_price("80", "1")
    assert detect_tools()["hardware_test_status"] == "NOT_RUN"
